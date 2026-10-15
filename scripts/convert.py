"""Convert the Minibar Excel workbook (Item Register sheet) to data/register.json."""
import json, sys, glob, datetime as dt
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else glob.glob("*.xlsx")[0]
wb = load_workbook(src, read_only=True, data_only=True)
ws = wb["Item Register"]

def d(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10] if v else None

rows = []
for r in ws.iter_rows(min_row=4, max_col=9, values_only=True):
    if not r[0] and not r[2]:
        continue
    try:
        qty = float(r[5] or 0)
        qty = int(qty) if qty == int(qty) else qty
    except Exception:
        qty = 0
    rows.append({"status": (r[0] or "").strip(), "type": (r[1] or "").strip(),
                 "item": (r[2] or "").strip(), "floor": str(r[3] or "").strip(),
                 "location": str(r[4] or "").strip(), "qty": qty,
                 "exp": d(r[6]), "condition": (r[7] or "").strip(), "tx": d(r[8])})

out = {"generated": dt.datetime.utcnow().isoformat(timespec="seconds") + "Z",
       "source": src.split("/")[-1], "rows": rows}
with open("data/register.json", "w") as f:
    json.dump(out, f, separators=(",", ":"))
print(f"{len(rows)} rows -> data/register.json")
